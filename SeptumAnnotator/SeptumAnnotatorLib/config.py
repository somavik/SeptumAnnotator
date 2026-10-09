"""Case configuration: the grind-mode table (CSV/XLSX) and single custom cases."""

from __future__ import annotations

import csv
import os
from dataclasses import asdict, dataclass, fields
from typing import Dict, List

# Structure name -> label value. VERIFY these against your data before annotating.
LABEL_MAPS: Dict[str, Dict[str, int]] = {
    # ImageCHD ground truth (label 5 = myocardium).
    "imagechd": {"lv": 1, "rv": 2, "la": 3, "ra": 4, "myo": 5, "aorta": 6, "pa": 7},
    # TotalSegmentator "heartchambers_highres" task.
    "totalseg_heartchambers": {"myo": 1, "la": 2, "lv": 3, "ra": 4, "rv": 5, "aorta": 6, "pa": 7},
}

SEG_SOURCES = ("gt", "generated")
SEED_STRATEGIES = ("myocardium", "sheetness")  # Plan A, Plan B
REGIONS = ("band", "hull")
OUTFLOWS = ("lv_aorta", "rv_pa")
# Ventricle and artery label names for each outflow.
OUTFLOW_STRUCTURES = {"lv_aorta": ("lv", "aorta"), "rv_pa": ("rv", "pa")}

DEFAULTS = {
    "spacing_mm": 1.0,
    "d_mm": 5.0,
    "sheetness_threshold": 0.1,
}

REQUIRED_COLUMNS = ("case_id", "ct_path", "seg_path")


class ConfigError(ValueError):
    pass


@dataclass
class CaseConfig:
    case_id: str
    ct_path: str
    seg_path: str
    seg_source: str = "gt"
    label_map: str = "imagechd"
    seed_strategy: str = "myocardium"
    region: str = "band"
    sheetness_path: str = ""
    skeleton_lv_aorta: str = ""
    skeleton_rv_pa: str = ""
    mitral_path: str = ""  # optional, display only (visual aid, protocol §9.2)

    def validate(self) -> None:
        problems = []
        if not self.case_id:
            problems.append("case_id is empty")
        for name in ("ct_path", "seg_path"):
            path = getattr(self, name)
            if not path or not os.path.isfile(path):
                problems.append(f"{name} not found: {path!r}")
        if self.seg_source not in SEG_SOURCES:
            problems.append(f"seg_source must be one of {SEG_SOURCES}, got {self.seg_source!r}")
        if self.label_map not in LABEL_MAPS:
            problems.append(f"label_map must be one of {tuple(LABEL_MAPS)}, got {self.label_map!r}")
        if self.seed_strategy not in SEED_STRATEGIES:
            problems.append(f"seed_strategy must be one of {SEED_STRATEGIES}, got {self.seed_strategy!r}")
        if self.region not in REGIONS:
            problems.append(f"region must be one of {REGIONS}, got {self.region!r}")
        if self.seed_strategy == "sheetness" and not os.path.isfile(self.sheetness_path):
            problems.append(f"seed_strategy 'sheetness' needs an existing sheetness_path, got {self.sheetness_path!r}")
        for outflow in OUTFLOWS:
            path = self.skeleton_path(outflow)
            if path and not os.path.isfile(path):
                problems.append(f"skeleton_{outflow} not found: {path!r}")
        if self.mitral_path and not os.path.isfile(self.mitral_path):
            problems.append(f"mitral_path not found: {self.mitral_path!r}")
        if problems:
            raise ConfigError(f"Case {self.case_id!r}: " + "; ".join(problems))

    def labels(self) -> Dict[str, int]:
        return LABEL_MAPS[self.label_map]

    def skeleton_path(self, outflow: str) -> str:
        return getattr(self, f"skeleton_{outflow}")

    def to_dict(self) -> dict:
        return asdict(self)


_FIELD_NAMES = [f.name for f in fields(CaseConfig)]
_PATH_FIELDS = ("ct_path", "seg_path", "sheetness_path", "skeleton_lv_aorta", "skeleton_rv_pa", "mitral_path")


def case_from_row(row: Dict[str, str], base_dir: str = "") -> CaseConfig:
    """Build a CaseConfig from one table row. Unknown columns are ignored on purpose:
    never put diagnosis information into the config (protocol §2, independence rule 1)."""
    values = {}
    for name in _FIELD_NAMES:
        raw = row.get(name)
        value = "" if raw is None else str(raw).strip()
        if value and name in _PATH_FIELDS and base_dir and not os.path.isabs(value):
            value = os.path.normpath(os.path.join(base_dir, value))
        if value:
            values[name] = value
    for name in REQUIRED_COLUMNS:
        values.setdefault(name, "")
    return CaseConfig(**values)


def _read_rows(path: str) -> List[Dict[str, str]]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
            reader = csv.DictReader(f, dialect=dialect)
            return [{(k or "").strip().lower(): v for k, v in r.items()} for r in reader]
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Reading .xlsx configs needs openpyxl (pip install openpyxl)") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:]]
    raise ConfigError(f"Unsupported config format {ext!r} (use .csv or .xlsx)")


def load_config(path: str, validate: bool = True) -> List[CaseConfig]:
    """Read the grind-mode table. Relative paths are resolved against the table's folder."""
    rows = _read_rows(path)
    if rows:
        missing = [c for c in REQUIRED_COLUMNS if c not in rows[0]]
        if missing:
            raise ConfigError(f"{path}: missing required columns {missing}")
    base_dir = os.path.dirname(os.path.abspath(path))
    cases, seen = [], set()
    for row in rows:
        if not any(str(v or "").strip() for v in row.values()):
            continue  # empty line
        case = case_from_row(row, base_dir)
        if case.case_id in seen:
            raise ConfigError(f"{path}: duplicate case_id {case.case_id!r}")
        seen.add(case.case_id)
        if validate:
            case.validate()
        cases.append(case)
    return cases
